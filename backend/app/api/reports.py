import io
import csv
from datetime import datetime, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.db.database import get_db
from app.models import AttendanceRecord, AttendanceSession, Student, ClassModel, AuditLog, User
from app.schemas import AttendanceReportSummary, AuditLogResponse
from app.api.auth import require_teacher_or_admin, require_admin

router = APIRouter(prefix="/reports", tags=["reports"])

@router.get("/attendance", response_model=AttendanceReportSummary)
def get_attendance_report(
    class_id: Optional[str] = None,
    section: Optional[str] = None,
    days: int = 14,
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher_or_admin)
):
    query = db.query(AttendanceRecord).join(AttendanceSession)

    if class_id and class_id != "all":
        query = query.filter(AttendanceSession.class_id == class_id)

    records = query.all()

    present_count = sum(1 for r in records if r.status == "present")
    late_count = sum(1 for r in records if r.status == "late")
    absent_count = sum(1 for r in records if r.status == "absent")
    total_records = len(records)

    attendance_rate = round(((present_count + late_count) / total_records * 100.0), 1) if total_records > 0 else 0.0

    today_str = datetime.utcnow().strftime("%Y-%m-%d")
    classes_today = db.query(AttendanceSession).filter(AttendanceSession.date == today_str).count()

    # Calculate 14-day trend
    trend = []
    end_date = datetime.utcnow()
    for i in range(days - 1, -1, -1):
        day = end_date - timedelta(days=i)
        day_str = day.strftime("%Y-%m-%d")
        day_label = day.strftime("%b %d")

        day_records = [r for r in records if r.session and r.session.date == day_str]
        d_present = sum(1 for r in day_records if r.status in ["present", "late"])
        d_total = len(day_records)
        rate = round((d_present / d_total * 100.0), 1) if d_total > 0 else 0.0

        trend.append({
            "date": day_str,
            "label": day_label,
            "rate": rate,
            "present": d_present,
            "absent": (d_total - d_present) if d_total > 0 else 0
        })

    # Class comparison
    classes = db.query(ClassModel).all()
    class_comparisons = []
    for c in classes:
        c_records = [r for r in records if r.session and r.session.class_id == c.id]
        c_pres = sum(1 for r in c_records if r.status in ["present", "late"])
        c_tot = len(c_records)
        c_rate = round((c_pres / c_tot * 100.0), 1) if c_tot > 0 else 0.0
        class_comparisons.append({
            "class_id": c.id,
            "name": f"{c.name} - {c.section}",
            "student_count": len(c.students),
            "attendance_rate": c_rate
        })

    return {
        "attendance_rate": attendance_rate,
        "present_count": present_count,
        "absent_count": absent_count,
        "late_count": late_count,
        "total_records": total_records,
        "classes_today": classes_today,
        "trend": trend,
        "class_comparisons": class_comparisons
    }

@router.get("/export-csv")
def export_attendance_csv(
    class_id: Optional[str] = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher_or_admin)
):
    """
    Generates downloadable CSV report for student attendance.
    """
    query = db.query(Student)
    if class_id and class_id != "all":
        classroom = db.query(ClassModel).filter(ClassModel.id == class_id).first()
        if classroom:
            query = query.filter(Student.classes.contains(classroom))

    students = query.order_by(Student.roll_number.asc()).all()

    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "Roll Number",
        "Student Name",
        "Class",
        "Section",
        "Total Sessions",
        "Present Count",
        "Late Count",
        "Absent Count",
        "Attendance Rate (%)",
        "Status"
    ])

    for s in students:
        records = s.attendance_records
        total = len(records)
        pres = sum(1 for r in records if r.status == "present")
        late = sum(1 for r in records if r.status == "late")
        abs = sum(1 for r in records if r.status == "absent")
        rate = f"{round(((pres + late) / total * 100.0), 1)}%" if total > 0 else "0.0%"

        writer.writerow([
            s.roll_number,
            s.name,
            s.class_name,
            s.section,
            total,
            pres,
            late,
            abs,
            rate,
            s.status
        ])

    output.seek(0)
    filename = f"attendai_report_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/export-excel")
def export_attendance_excel(
    class_id: Optional[str] = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_teacher_or_admin)
):
    """
    Generates professionally formatted Excel (.xlsx) workbook using openpyxl.
    """
    query = db.query(Student)
    if class_id and class_id != "all":
        classroom = db.query(ClassModel).filter(ClassModel.id == class_id).first()
        if classroom:
            query = query.filter(Student.classes.contains(classroom))

    students = query.order_by(Student.roll_number.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"

    # Title & Metadata
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "Kanpur Institute of Technology (KIT Kanpur) — Attendance Management Report"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="171717")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Generated on {datetime.utcnow().strftime('%B %d, %Y %H:%M UTC')} | Requested by: {user.full_name}"
    sub_cell.font = Font(name="Arial", size=10, italic=True, color="737373")

    headers = [
        "Roll Number", "Student Name", "Class", "Section",
        "Total Sessions", "Present", "Late", "Absent", "Rate (%)", "Status"
    ]

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3157D5", end_color="3157D5", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    ws.append([])  # Row 3 empty
    ws.append(headers)  # Row 4

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=4, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align

    thin_border = Border(
        left=Side(style="thin", color="E7E7E3"),
        right=Side(style="thin", color="E7E7E3"),
        top=Side(style="thin", color="E7E7E3"),
        bottom=Side(style="thin", color="E7E7E3"),
    )

    for s in students:
        records = s.attendance_records
        total = len(records)
        pres = sum(1 for r in records if r.status == "present")
        late = sum(1 for r in records if r.status == "late")
        abs_cnt = sum(1 for r in records if r.status == "absent")
        rate = round(((pres + late) / total * 100.0), 1) if total > 0 else 0.0

        row = [
            s.roll_number,
            s.name,
            s.class_name,
            s.section,
            total,
            pres,
            late,
            abs_cnt,
            f"{rate}%",
            s.status.capitalize()
        ]
        ws.append(row)
        current_row = ws.max_row
        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            if col_num in [1, 4, 5, 6, 7, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="center")

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    filename = f"attendai_attendance_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/audit")
def get_audit_logs(
    limit: int = 50,
    db: Session = Depends(get_db),
    user: User = Depends(require_admin)
):
    """
    Returns audit trail of manual adjustments. Restricted to Administrators.
    """
    logs = db.query(AuditLog).order_by(AuditLog.timestamp.desc()).limit(limit).all()
    results = []
    for l in logs:
        results.append({
            "id": l.id,
            "user_name": l.user.full_name if l.user else "System",
            "action": l.action,
            "entity_type": l.entity_type,
            "entity_id": l.entity_id,
            "details": l.details,
            "timestamp": l.timestamp
        })
    return results
